import os
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
import docx

os.makedirs("demo_documents", exist_ok=True)

# 1. Generate PDF: SOP-QA-042-Deviation-Management.pdf
pdf_path = os.path.join("demo_documents", "SOP-QA-042-Deviation-Management.pdf")
doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
styles = getSampleStyleSheet()

title_style = ParagraphStyle(
    "DocTitle",
    parent=styles["Title"],
    fontSize=18,
    leading=22,
    textColor=colors.HexColor("#0f172a"),
    alignment=0,
)
h1_style = ParagraphStyle(
    "Heading1",
    parent=styles["Heading1"],
    fontSize=13,
    leading=17,
    textColor=colors.HexColor("#1e293b"),
    spaceBefore=12,
    spaceAfter=6,
)
body_style = ParagraphStyle(
    "Body",
    parent=styles["Normal"],
    fontSize=10,
    leading=14,
    textColor=colors.HexColor("#334155"),
    spaceAfter=8,
)

story = []

# Page 1
story.append(Paragraph("ACME PHARMA — STANDARD OPERATING PROCEDURE", title_style))
story.append(Paragraph("Document ID: SOP-QA-042 | Version: 4.0 | Effective Date: 2026-01-15", body_style))
story.append(Paragraph("Title: Deviation Management & Root Cause Analysis Procedure", h1_style))
story.append(Spacer(1, 10))

story.append(Paragraph("1. Purpose and Scope", h1_style))
story.append(Paragraph(
    "This Standard Operating Procedure (SOP) defines the mandatory protocol for identifying, documenting, "
    "evaluating, and resolving planned and unplanned deviations in Good Manufacturing Practice (GMP) operations "
    "across all Acme Pharma production, packaging, and quality control facilities. It applies to all personnel, "
    "contractors, and laboratory analysts involved in commercial pharmaceutical manufacturing.",
    body_style,
))

story.append(Paragraph("2. Regulatory References & Compliance", h1_style))
story.append(Paragraph(
    "This procedure strictly adheres to US FDA 21 CFR 211.192 (Production record review), EU GMP Guidelines Volume 4 "
    "Part 1 Chapter 1 (Pharmaceutical Quality System), and ICH Q10 Pharmaceutical Quality System. All deviation "
    "investigations must be documented contemporaneously within the validated Enterprise QMS platform.",
    body_style,
))

story.append(Paragraph("3. Deviation Classification", h1_style))
story.append(Paragraph(
    "Deviations are categorized into three severity tiers based on risk to patient safety, product identity, and purity:<br/>"
    "• <b>Critical Deviation:</b> Departure directly impacting product safety, potency, or efficacy. Requires QA notification within 2 hours.<br/>"
    "• <b>Major Deviation:</b> Non-compliance with validated manufacturing limits or analytical specifications without direct safety compromise.<br/>"
    "• <b>Minor Deviation:</b> Minor clerical or non-critical process variations having zero impact on product quality or yield.",
    body_style,
))

story.append(PageBreak())

# Page 2
story.append(Paragraph("4. Immediate Containment & Batch Quarantine", h1_style))
story.append(Paragraph(
    "Upon discovering a deviation, the discovering operator must immediately initiate containment actions. "
    "The affected batch, lot, or equipment must be placed in electronically locked quarantine status within the ERP system. "
    "Physical warning tags ('BATCH QUARANTINE - DO NOT MOVE') must be physically affixed to all affected material pallets "
    "within 30 minutes of discovery.",
    body_style,
))

story.append(Paragraph("5. Root Cause Analysis Methodology", h1_style))
story.append(Paragraph(
    "All Major and Critical deviations require structured Root Cause Analysis (RCA). Permitted RCA tools include the "
    "Ishikawa Fishbone diagram, 5 Whys, and Failure Mode and Effects Analysis (FMEA). Human error may never be cited as "
    "the standalone root cause without an evaluation of procedural clarity, training efficacy, and workstation ergonomics.",
    body_style,
))

story.append(Paragraph("6. Corrective and Preventive Action (CAPA) Initiation", h1_style))
story.append(Paragraph(
    "Where the investigation reveals systemic or recurring risk, a formal CAPA record must be linked to the deviation. "
    "CAPA plans must specify: measurable action items, designated department owners, verification criteria, and a 90-day "
    "post-implementation effectiveness review date.",
    body_style,
))

story.append(PageBreak())

# Page 3
story.append(Paragraph("7. Investigation Timelines & Closure SLAs", h1_style))
story.append(Paragraph(
    "Investigation closure must adhere to rigorous service level agreements from the date of initial discovery:<br/>"
    "• <b>Critical Deviations:</b> Comprehensive report and QA sign-off within 5 business days.<br/>"
    "• <b>Major Deviations:</b> Full investigation and CAPA assignment within 15 business days.<br/>"
    "• <b>Minor Deviations:</b> Batch documentation closeout within 30 calendar days.<br/>"
    "Extensions require formal written approval from the Head of Quality Assurance prior to SLA breach.",
    body_style,
))

story.append(Paragraph("8. Records Archival and Annual Review", h1_style))
story.append(Paragraph(
    "All completed deviation records, analytical chromatograms, and witness statements must be retained for a minimum of "
    "7 years or 1 year past product expiry date (whichever is greater). Deviations are reviewed quarterly by the Quality "
    "Management Board to identify repetitive failure modes.",
    body_style,
))

doc.build(story)
print(f"Created: {pdf_path}")

# 2. Generate DOCX: HR-POL-108-Leave-and-Attendance-Policy.docx
docx_path = os.path.join("demo_documents", "HR-POL-108-Leave-and-Attendance-Policy.docx")
d = docx.Document()
d.add_heading("ACME PHARMA — HUMAN RESOURCES POLICY", 0)
d.add_paragraph("Policy Code: HR-POL-108 | Effective Date: 2026-01-01 | Target Audience: All Global Employees")

d.add_heading("1. Purpose & Policy Scope", level=1)
d.add_paragraph(
    "This policy defines annual paid time off, sick leave, compassionate leave, and parental benefits for all full-time "
    "and part-time staff members at Acme Pharma. We are committed to fostering employee wellness and balanced workplace productivity."
)

d.add_heading("2. Annual Paid Time Off (PTO) & Accrual Limits", level=1)
d.add_paragraph(
    "Full-time employees accrue 20 days of paid time off per calendar year, accrued semi-monthly. "
    "A maximum of 5 unused PTO days may be rolled over into the subsequent calendar year, expiring on March 31st. "
    "PTO requests exceeding 3 consecutive business days must be submitted via the HR Portal at least 14 days in advance."
)

d.add_heading("3. Parental and Family Care Leave", level=1)
d.add_paragraph(
    "Acme Pharma provides 16 weeks of fully paid parental leave for primary caregivers and 6 weeks for secondary caregivers, "
    "applicable to biological births, adoptions, and foster placements. Leave can be taken continuously or in two separate blocks "
    "within the first 12 months."
)

d.add_heading("4. Sick Leave and Medical Certification", level=1)
d.add_paragraph(
    "Employees receive 10 days of paid sick leave annually on January 1st. Consecutive medical absences of 3 or more days "
    "require a certified physician note submitted directly to the confidential HR Medical Benefits mailbox within 48 hours of return."
)

d.save(docx_path)
print(f"Created: {docx_path}")

# 3. Generate TXT: IT-SEC-015-Password-and-Access-Control-Policy.txt
txt_path = os.path.join("demo_documents", "IT-SEC-015-Password-and-Access-Control-Policy.txt")
txt_content = """ACME PHARMA GLOBAL INFORMATION SECURITY POLICY
Document Ref: IT-SEC-015
Classification: INTERNAL CONFIDENTIAL
Effective Date: 2026-02-01

SECTION 1: SCOPE AND APPLICABILITY
This policy governs all access credentials, identity lifecycle events, and workstation authentication mechanisms across all corporate workstations, manufacturing network domains (SCADA/LIMS), and cloud infrastructure at Acme Pharma.

SECTION 2: PASSWORD COMPLEXITY AND LIFECYCLE
1. Minimum length: Passwords must contain a minimum of 14 characters.
2. Character variety: Must include at least 1 uppercase letter, 1 lowercase letter, 1 numeric digit, and 1 approved special character (@, #, $, %, ^, &, *).
3. Expiration: Passwords must be rotated every 90 days. System prevents reuse of the last 12 historical passwords.
4. Account Lockout: Workstations and portal logins lock automatically after 5 consecutive failed attempts. Unlock requires IT Helpdesk identity verification.

SECTION 3: MULTI-FACTOR AUTHENTICATION (MFA)
Multi-Factor Authentication (MFA) via FIDO2 hardware keys or approved mobile authenticator push notifications is mandatory for:
- All remote VPN and corporate portal logins.
- Any access to production databases containing customer, clinical trial, or batch records.
- Administrative access to AWS and Supabase infrastructure.
SMS-based verification is explicitly prohibited due to SIM-swap vulnerabilities.

SECTION 4: WORKSTATION AND SCREEN LOCK TIMEOUTS
All unattended employee workstations must lock automatically after 5 minutes of inactivity. Screen privacy filters are mandatory in manufacturing and laboratory common areas.
"""
with open(txt_path, "w", encoding="utf-8") as f:
    f.write(txt_content.strip())
print(f"Created: {txt_path}")

# 4. Generate MD: CLIN-SOP-009-Cold-Chain-Storage.md
md_path = os.path.join("demo_documents", "CLIN-SOP-009-Cold-Chain-Storage.md")
md_content = """# CLINICAL OPERATIONS SOP: REFRIGERATED STORAGE & COLD CHAIN INTEGRITY

**Document ID:** CLIN-SOP-009  
**Facility:** Acme Central Logistics & Regional Depots  
**Standard:** GDP / USP <659> Packaging and Storage Requirements  

## 1. Temperature Specifications and Monitoring
All temperature-sensitive biopharmaceuticals, vaccine candidates, and clinical trial samples must be maintained within the following strict bands:
- **Refrigerated Storage:** Controlled between +2.0°C and +8.0°C (+35.6°F to +46.4°F).
- **Deep Frozen:** Maintained at -20.0°C ± 5°C.
- **Ultra-Low Storage:** Maintained between -80.0°C and -65.0°C.

Automated IoT temperature sensors record ambient readings at 5-minute intervals. Calibration of all sensor probes is certified semi-annually against NIST-traceable standards.

## 2. Temperature Excursion Management
An excursion is defined as any verified temperature deviation exceeding allowable bands for greater than 15 continuous minutes:
1. Immediate audio and visual alarms trigger across the central dispatch console.
2. The on-call Cold Chain Engineer is automatically notified via priority paging.
3. If primary cooling fails to recover within 30 minutes, samples must be immediately transferred to the redundant backup refrigeration bay (Unit B-12).

## 3. Quarantine and Disposition Protocol
Any batch subjected to an unapproved excursion must be tagged in 'TEMP QUARANTINE' status within the LIMS system. Stability testing and Quality Assurance sign-off are required prior to releasing materials back to clinical trial distribution.
"""
with open(md_path, "w", encoding="utf-8") as f:
    f.write(md_content.strip())
print(f"Created: {md_path}")
